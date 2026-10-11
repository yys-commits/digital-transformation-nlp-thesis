# -*- coding: utf-8 -*-
"""
把论文 Markdown 转换为 Word（.docx）
支持：多级标题、markdown 表格、图片、引用、列表、加粗与行内代码。
用法：python3 tools/md_to_docx.py
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "paper" / "毕业论文_企业数字化转型文本测度研究.md"
DST = ROOT / "paper" / "毕业论文_企业数字化转型文本测度研究.docx"

CN_FONT = "宋体"
EN_FONT = "Times New Roman"


def set_font(run, size=None, bold=None, name=CN_FONT):
    run.font.name = EN_FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_rich_text(par, text, size=11, base_bold=False):
    """处理 **加粗** 与 `代码`"""
    import re
    for seg in re.split(r"(\*\*.+?\*\*|`.+?`)", text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = par.add_run(seg[2:-2]); set_font(r, size, True)
        elif seg.startswith("`") and seg.endswith("`"):
            r = par.add_run(seg[1:-1]); set_font(r, size, base_bold); r.font.color.rgb = RGBColor(0xA0, 0x30, 0x30)
        else:
            r = par.add_run(seg); set_font(r, size, base_bold)


def is_table_sep(line):
    s = line.strip()
    return s.startswith("|") and set(s.replace("|", "").replace(" ", "")) <= set("-:")


def parse_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    style.font.size = Pt(11)

    lines = SRC.read_text(encoding="utf-8").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()

        # 表格
        if s.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            header = parse_row(lines[i])
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(parse_row(lines[i]))
                i += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Table Grid"
            for j, h in enumerate(header):
                cell = t.rows[0].cells[j]
                cell.text = ""
                add_rich_text(cell.paragraphs[0], h, size=10, base_bold=True)
            for r in rows:
                cells = t.add_row().cells
                for j, v in enumerate(r[:len(header)]):
                    cells[j].text = ""
                    add_rich_text(cells[j].paragraphs[0], v, size=10)
            doc.add_paragraph()
            continue

        # 图片
        if s.startswith("!["):
            path = s[s.index("(") + 1:s.rindex(")")]
            img = (ROOT / "paper" / Path(path)).resolve()
            if img.exists():
                doc.add_picture(str(img), width=Inches(5.8))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_font(cap.add_run(s[2:s.index("]")]), size=9)
            i += 1
            continue

        # 标题
        if s.startswith("#"):
            lvl = len(s) - len(s.lstrip("#"))
            h = doc.add_heading(level=min(lvl, 4))
            add_rich_text(h, s.lstrip("# ").strip(), size=max(17 - lvl * 2, 12), base_bold=True)
            i += 1
            continue

        # 分隔线
        if s in {"---", "***", "___"}:
            doc.add_paragraph("─" * 40)
            i += 1
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            add_rich_text(p, s.lstrip("> ").strip(), size=10.5)
            i += 1
            continue

        # 列表
        if s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, s[2:].strip(), size=11)
            i += 1
            continue

        # 空行
        if not s:
            i += 1
            continue

        # 正文
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        add_rich_text(p, s, size=11)
        i += 1

    doc.save(str(DST))
    print(f"已生成 Word 文档：{DST}（{DST.stat().st_size/1024:.0f} KB）")


if __name__ == "__main__":
    main()
